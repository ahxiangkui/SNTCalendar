#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
单位换算试题生成器 - 二年级数学练习题
Unit Conversion Test Generator - 2nd Grade Math Practice
"""

import random
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


# 配置常量
FILL_BLANK_RATIO = 0.7  # 填空题占比70%


class UnitConversionGenerator:
    """生成单位换算试题的类"""
    
    def __init__(self):
        """初始化生成器"""
        # 长度单位换算关系：1米 = 100厘米
        self.length_units = {
            '米': 100,
            '厘米': 1
        }
        
        # 时间单位换算关系
        self.time_units = {
            '时': 3600,  # 秒
            '分': 60,    # 秒
            '秒': 1
        }
        
        # 人民币单位换算关系：1元 = 10角
        self.currency_units = {
            '元': 10,
            '角': 1
        }
    
    def generate_length_conversion(self):
        """生成长度单位换算题"""
        # 随机选择换算方向
        if random.choice([True, False]):
            # 米转厘米
            meters = random.randint(1, 20)
            answer = meters * 100
            return f"{meters}米 = (    )厘米", answer
        else:
            # 厘米转米（选择能整除的数）
            meters = random.randint(1, 15)
            centimeters = meters * 100
            return f"{centimeters}厘米 = (    )米", meters
    
    def generate_time_conversion(self):
        """生成时间单位换算题"""
        conversion_type = random.choice(['hour_to_min', 'min_to_sec', 'hour_to_sec'])
        
        if conversion_type == 'hour_to_min':
            # 时转分钟
            hours = random.randint(1, 10)
            answer = hours * 60
            return f"{hours}时 = (    )分", answer
        elif conversion_type == 'min_to_sec':
            # 分钟转秒
            minutes = random.randint(1, 10)
            answer = minutes * 60
            return f"{minutes}分 = (    )秒", answer
        else:
            # 时转秒
            hours = random.randint(1, 5)
            answer = hours * 3600
            return f"{hours}时 = (    )秒", answer
    
    def generate_currency_conversion(self):
        """生成人民币单位换算题"""
        if random.choice([True, False]):
            # 元转角
            yuan = random.randint(1, 20)
            answer = yuan * 10
            return f"{yuan}元 = (    )角", answer
        else:
            # 角转元（选择能整除的数）
            yuan = random.randint(1, 15)
            jiao = yuan * 10
            return f"{jiao}角 = (    )元", yuan
    
    def generate_comparison_question(self):
        """生成比大小题目"""
        question_type = random.choice(['length', 'time', 'currency'])
        
        if question_type == 'length':
            # 长度比较
            if random.choice([True, False]):
                # 米和厘米比较
                meters = random.randint(1, 10)
                # 生成50-249之间的厘米数，避免太简单的比较
                centimeters = random.randint(50, 249)
                left = f"{meters}米"
                right = f"{centimeters}厘米"
                left_value = meters * 100
                right_value = centimeters
            else:
                # 两个相同单位比较
                if random.choice([True, False]):
                    # 都是米
                    val1 = random.randint(1, 50)
                    val2 = random.randint(1, 50)
                    left = f"{val1}米"
                    right = f"{val2}米"
                    left_value = val1
                    right_value = val2
                else:
                    # 都是厘米
                    val1 = random.randint(50, 500)
                    val2 = random.randint(50, 500)
                    left = f"{val1}厘米"
                    right = f"{val2}厘米"
                    left_value = val1
                    right_value = val2
        
        elif question_type == 'time':
            # 时间比较
            conversions = [
                ('时', '分', 60),
                ('分', '秒', 60)
            ]
            unit1, unit2, factor = random.choice(conversions)
            
            if random.choice([True, False]):
                # 不同单位
                val1 = random.randint(1, 5)
                val2 = random.randint(60, 400)
                left = f"{val1}{unit1}"
                right = f"{val2}{unit2}"
                left_value = val1 * factor
                right_value = val2
            else:
                # 相同单位
                val1 = random.randint(10, 100)
                val2 = random.randint(10, 100)
                left = f"{val1}{unit2}"
                right = f"{val2}{unit2}"
                left_value = val1
                right_value = val2
        
        else:  # currency
            # 人民币比较
            if random.choice([True, False]):
                # 元和角比较
                yuan = random.randint(1, 10)
                jiao = random.randint(5, 150)
                left = f"{yuan}元"
                right = f"{jiao}角"
                left_value = yuan * 10
                right_value = jiao
            else:
                # 相同单位
                if random.choice([True, False]):
                    val1 = random.randint(1, 20)
                    val2 = random.randint(1, 20)
                    left = f"{val1}元"
                    right = f"{val2}元"
                    left_value = val1
                    right_value = val2
                else:
                    val1 = random.randint(5, 100)
                    val2 = random.randint(5, 100)
                    left = f"{val1}角"
                    right = f"{val2}角"
                    left_value = val1
                    right_value = val2
        
        # 确定比较符号
        if left_value > right_value:
            answer = '>'
        elif left_value < right_value:
            answer = '<'
        else:
            answer = '='
        
        return f"{left} ○ {right}", answer
    
    def generate_fill_in_blank_question(self):
        """生成填空题"""
        question_type = random.choice(['length', 'time', 'currency'])
        
        if question_type == 'length':
            return self.generate_length_conversion()
        elif question_type == 'time':
            return self.generate_time_conversion()
        else:
            return self.generate_currency_conversion()


def create_test_paper(filename="二年级单位换算试题.docx", pages=10, questions_per_page=50):
    """
    创建单位换算试题Word文档
    
    Args:
        filename: 输出文件名
        pages: 页数
        questions_per_page: 每页题目数
    """
    doc = Document()
    
    # 设置中文字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.styles['Normal'].font.size = Pt(10.5)
    
    generator = UnitConversionGenerator()
    
    for page_num in range(1, pages + 1):
        # 添加标题
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(f'二年级数学单位换算练习题（第{page_num}页）')
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.name = '黑体'
        title_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        
        # 添加说明
        instruction = doc.add_paragraph()
        instruction.alignment = WD_ALIGN_PARAGRAPH.CENTER
        instruction_run = instruction.add_run('姓名：________  班级：________  日期：________  得分：________')
        instruction_run.font.size = Pt(10.5)
        
        doc.add_paragraph()  # 空行
        
        # 生成题目
        # 填空题部分（约70%）
        section1 = doc.add_paragraph()
        section1_run = section1.add_run('一、单位换算填空题（每题2分）')
        section1_run.font.size = Pt(12)
        section1_run.font.bold = True
        
        fill_blank_count = int(questions_per_page * FILL_BLANK_RATIO)
        comparison_count = questions_per_page - fill_blank_count
        
        for i in range(1, fill_blank_count + 1):
            question, answer = generator.generate_fill_in_blank_question()
            p = doc.add_paragraph(style='List Number')
            p_run = p.add_run(question)
            p_run.font.size = Pt(10.5)
            
            # 每5题空一行
            if i % 5 == 0 and i < fill_blank_count:
                doc.add_paragraph()
        
        doc.add_paragraph()  # 空行
        
        # 比大小部分（约15题）
        section2 = doc.add_paragraph()
        section2_run = section2.add_run('二、比大小（在○里填上">"、"<"或"="）（每题2分）')
        section2_run.font.size = Pt(12)
        section2_run.font.bold = True
        
        for i in range(1, comparison_count + 1):
            question, answer = generator.generate_comparison_question()
            p = doc.add_paragraph(style='List Number')
            p_run = p.add_run(question)
            p_run.font.size = Pt(10.5)
            
            # 每5题空一行
            if i % 5 == 0 and i < comparison_count:
                doc.add_paragraph()
        
        # 如果不是最后一页，添加分页符
        if page_num < pages:
            doc.add_page_break()
    
    # 保存文档
    doc.save(filename)
    print(f"试题已生成：{filename}")
    print(f"共{pages}页，每页约{questions_per_page}题")


if __name__ == "__main__":
    # 生成试题
    create_test_paper(
        filename="二年级单位换算试题.docx",
        pages=10,
        questions_per_page=50
    )
